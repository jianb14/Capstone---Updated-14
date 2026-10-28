"""Exports preserve filter semantics and cover every displayed report series."""
import io
import os
from datetime import datetime, time, timezone as utc
from pathlib import Path
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

import openpyxl
from django.conf import settings
from django.test import TestCase, override_settings
from django.urls import reverse
from pypdf import PdfReader

from .analytics_reporting import build_report, resolve_filters
from .models import Booking, Payment, Review, User


@override_settings(REQUEST_RATE_LIMITING_ENABLED=False)
class AnalyticsExportParityTests(TestCase):
    def setUp(self):
        self.admin=User.objects.create_user(username='export-admin',role='admin',email_verified=True)
        self.customer=User.objects.create_user(username='=FORMULA()',role='customer',email_verified=True)
        self.client.force_login(self.admin)
        self.clock=self.enterContext(patch('app.analytics_reporting.timezone.now',return_value=datetime(2026,10,26,4,tzinfo=utc.utc)))
        self.done=self.booking('2026-10-26','completed','1250.25')
        self.previous=self.booking('2026-09-26','completed','500')
        self.booking('2026-09-27','completed','900')
        self.booking('2026-10-26','confirmed','2000',event_type='Wedding')
        self.payment(self.done,'750.15',datetime(2026,10,26,2,tzinfo=utc.utc))
        self.payment(self.previous,'100',datetime(2026,9,26,2,tzinfo=utc.utc))
        Review.objects.create(booking=self.done,user=self.customer,rating=4,comment='Good styling')

    def booking(self,day,status,value,**kwargs):
        return Booking.objects.create(user=self.customer,event_date=day,event_time=time(9),status=status,total_price=value,
            event_location='Test venue',package_type='Package 1-A + Entrance A + Stand A',event_type=kwargs.get('event_type','Birthday'))

    def payment(self,booking,value,stamp):
        return Payment.objects.create(booking=booking,amount=value,payment_status='verified',payment_method='gcash',
            received_at=stamp,transaction_ref=f'EXPORT-{booking.pk}')

    def report(self,**kwargs):
        return build_report(resolve_filters({'filter_preset':'month','event_type':'Birthday',**kwargs}))

    def test_frozen_month_export_retains_matching_previous_month_even_after_midnight(self):
        report=self.report()
        query={key:values[0] for key,values in parse_qs(urlsplit(report['export_url']).query).items()}
        self.assertEqual(query['comparison'],'month')
        self.assertNotIn('filter_preset',query)
        self.clock.return_value=datetime(2026,11,2,4,tzinfo=utc.utc)
        frozen=build_report(resolve_filters(query))
        self.assertEqual(frozen['filters'].period_label,report['filters'].period_label)
        self.assertEqual(frozen['filters'].previous_start,report['filters'].previous_start)
        self.assertEqual(frozen['filters'].previous_end,report['filters'].previous_end)
        self.assertEqual(frozen['previous_count'],1)
        self.assertEqual(frozen['comparison'],report['comparison'])
        self.assertEqual(frozen['insights']['collections']['trend'],report['insights']['collections']['trend'])

    def test_all_time_export_keeps_comparison_disabled(self):
        report=self.report(filter_preset='all')
        query={key:values[0] for key,values in parse_qs(urlsplit(report['export_url']).query).items()}
        frozen=build_report(resolve_filters(query))
        self.assertIsNone(frozen['previous_count'])
        self.assertFalse(frozen['charts']['trend']['has_comparison'])

    def test_pdf_and_excel_links_share_every_applied_filter_from_any_tab(self):
        for tab in ('overview','bookings','payments','customers'):
            response=self.client.get(reverse('admin_analytics'),{'filter_preset':'month','event_type':'Birthday','tab':tab},HTTP_X_REQUESTED_WITH='XMLHttpRequest')
            payload=response.json()
            self.assertEqual(parse_qs(urlsplit(payload['export_url']).query),parse_qs(urlsplit(payload['excel_url']).query))
            self.assertEqual(parse_qs(urlsplit(payload['excel_url']).query)['event_type'],['Birthday'])

    def test_excel_covers_all_chart_series_and_preserves_numeric_cells_and_filter_context(self):
        long_comment = 'Full review content. ' * 100
        Review.objects.filter(booking=self.done).update(comment=long_comment,
            created_at=datetime(2026,10,25,23,30,tzinfo=utc.utc))
        report=self.report()
        response=self.client.get(report['excel_url'])
        self.assertEqual(response.status_code,200)
        workbook=openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook.sheetnames[0],'Overview')
        self.assertEqual(workbook['Bookings Summary'].max_row,2)
        self.assertEqual(workbook['Booking Trends'].max_row,27)
        trend=list(workbook['Booking Trends'].values)[1:]
        self.assertEqual(sum(row[1] for row in trend),report['total_bookings'])
        self.assertEqual(sum(row[5] or 0 for row in trend),report['previous_count'])
        self.assertAlmostEqual(sum(row[2] for row in trend),float(report['completed_value']))
        self.assertAlmostEqual(sum(row[3] for row in trend),float(report['verified_collections']))
        receipts=list(workbook['Collection Trends'].values)[1:]
        self.assertAlmostEqual(sum(row[1] for row in receipts),float(report['insights']['collections']['value']))
        for name in ('Event Demand','Offer Demand','Payment Breakdown','Customer Breakdown','Booking Breakdown'):
            self.assertGreater(workbook[name].max_row,1)
        metadata=dict(workbook['Report Definitions'].values)
        self.assertEqual(metadata['Comparison'],report['filters'].comparison_label)
        self.assertEqual(metadata['Event type'],'Birthday')
        for sheet in workbook:
            self.assertEqual(sheet.freeze_panes,'A2')
            self.assertTrue(sheet.auto_filter.ref)
            self.assertFalse(sheet.sheet_view.showGridLines)
        self.assertEqual(workbook['Bookings Summary']['B2'].value,'=FORMULA()')
        self.assertEqual(workbook['Bookings Summary']['B2'].data_type,'s')
        self.assertEqual(workbook['Booking Trends']['C2'].data_type,'n')
        self.assertEqual(workbook['Reviews']['D2'].value,long_comment)
        self.assertEqual(workbook['Reviews']['F2'].value,'2026-10-26 07:30')

    def test_real_pdf_includes_all_filtered_views_and_trend_appendix(self):
        report=self.report()
        response=self.client.get(report['export_url'])
        self.assertEqual(response.status_code,200)
        pdf=PdfReader(io.BytesIO(response.content))
        content='\n'.join(page.extract_text() or '' for page in pdf.pages)
        for label in ('Key Performance Indicators','Booking Outcomes and Demand','Payments and Balances',
                'Customers and Reviews','Review Rating Distribution','Collections by Payment Method',
                'Booking Trend / Selected Period','Booking Trend / Previous Period','Collections Received / Period Comparison'):
            self.assertIn(label,content)
        self.assertIn('Sep 26, 2026',content)
        self.assertIn('1,250.25',content)
        self.assertIn('750.15',content)
        self.assertIn('GCash',content)
        self.assertNotIn('Wedding',content)
        self.assertLess(len(pdf.pages),14)
        for page in pdf.pages:
            self.assertEqual(round(float(page.mediabox.width)),595)
            self.assertIn('Balloorina',page.extract_text())
        if os.environ.get('CAPTURE_ANALYTICS_EXPORT')=='1':
            folder=Path(settings.BASE_DIR)/'.booking-followup-check'
            (folder/'analytics-report-preview.pdf').write_bytes(response.content)
            (folder/'analytics-workbook-preview.xlsx').write_bytes(self.client.get(report['excel_url']).content)
