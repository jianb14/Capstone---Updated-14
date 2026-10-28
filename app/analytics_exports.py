"""Export every displayed analytics series from the shared reporting result."""
from decimal import Decimal
from math import ceil


def comparison_rows(report):
    return [
        {'label': 'Bookings', 'current': report['total_bookings'], 'previous': report['previous_count'], 'change': report['comparison']['count']['text'], 'money': False},
        {'label': 'Completed booking value', 'current': report['completed_value'], 'previous': report['previous_completed_value'], 'change': report['comparison']['value']['text'], 'money': True},
        {'label': 'Cohort collections', 'current': report['verified_collections'], 'previous': report['previous_verified_collections'], 'change': report['comparison']['collections']['text'], 'money': True},
        {'label': 'Collections received', 'current': report['insights']['collections']['value'], 'previous': report['insights']['collections']['previous_value'], 'change': report['insights']['collections']['delta']['text'], 'money': True},
    ]


def trend_rows(report):
    trend = report['charts']['trend']
    rows = []
    for index, period in enumerate(trend['ranges']):
        row = {'period': period, 'previous_period': trend['previous_ranges'][index] or 'No matching interval'}
        for key, series in trend['series'].items():
            row[key] = series['current'][index]
            row['previous_' + key] = series['previous'][index] if series['previous'] else None
        rows.append(row)
    return rows


def receipt_rows(report):
    trend = report['insights']['collections']['trend']
    return [{'period': period, 'value': trend['current'][index],
        'previous_period': trend['previous_ranges'][index] or 'No matching interval',
        'previous': trend['previous'][index] if trend['previous'] else None}
        for index, period in enumerate(trend['ranges'])]


def add_analytics_sheets(workbook, report):
    """Detail sheets supplement, rather than substitute for, the chart data."""
    overview = workbook.create_sheet('Overview', 0)
    overview.append(['Metric', 'Value', 'Definition'])
    for item in report['kpis']:
        key = item['key']
        value = {'bookings': report['total_bookings'], 'completed_value': report['completed_value'],
            'collections': report['verified_collections'], 'completion': report['completion_rate'],
            'cancellation': report['cancellation_rate'], 'customers': report['customer_count']}[key]
        overview.append([item['label'] + (' (%)' if key in ('completion', 'cancellation') else ''), value, item['note']])
    overview.append(['Average completed value (PHP)', report['average_completed_value'], 'Per completed booking'])
    overview.append(['Date basis', report['filters'].basis_label])
    overview.append(['Period', report['filters'].period_label])
    overview.append(['Comparison', report['filters'].comparison_label])
    overview.append(['Event type', report['filters'].event_type])
    overview.append(['Generated (Philippine time)', report['generated_at'].strftime('%Y-%m-%d %H:%M:%S')])
    comparisons = workbook.create_sheet('Period Comparisons')
    comparisons.append(['Metric', 'Selected period', 'Previous period', 'Change'])
    for row in comparison_rows(report):
        comparisons.append([row['label'], row['current'], row['previous'], row['change']])
    facts = workbook['Business Insights']
    for label, value in [('Committed events', report['insights']['committed_count']),
            ('Missing event start times', report['insights']['unknown_time']),
            ('Unclassified offer selections', report['insights']['offers']['unclassified']),
            ('Customers served', report['insights']['customers']['completed_customers']),
            ('Returning customer share (%)', report['insights']['customers']['returning_rate']),
            ('Eligible completed events', report['insights']['customers']['eligible_completed']),
            ('Reviewed completed events', report['insights']['customers']['reviewed_bookings'])]:
        facts.append([label, value])
    trend = workbook.create_sheet('Booking Trends')
    trend.append(['Selected dates', 'Bookings', 'Completed value (PHP)', 'Cohort collections (PHP)',
        'Previous dates', 'Previous bookings', 'Previous completed value (PHP)', 'Previous cohort collections (PHP)'])
    for row in trend_rows(report):
        trend.append([row[key] for key in ('period', 'bookings', 'completed_value', 'collections',
            'previous_period', 'previous_bookings', 'previous_completed_value', 'previous_collections')])
    breakdown = workbook.create_sheet('Booking Breakdown')
    breakdown.append(['Section', 'Label', 'Bookings', 'Share (%)', 'Completed value (PHP)', 'Cohort collections (PHP)'])
    for row in report['statuses']:
        breakdown.append(['Status', row['label'], row['count'], row['share']])
    for row in report['event_rows']:
        breakdown.append(['Event mix', row['label'], row['count'], None, row['value'], row['collections']])
    for row in report['setup_rows']:
        breakdown.append(['Main setup', row['label'], row['count']])
    demand = workbook.create_sheet('Event Demand')
    demand.append(['Grouping', 'Event period', 'Committed events'])
    for group in ('months', 'weekdays', 'hours'):
        for row in report['insights'][group]:
            demand.append([group.title(), row['label'], row['count']])
    receipts = workbook.create_sheet('Collection Trends')
    receipts.append(['Selected dates', 'Received collections (PHP)', 'Previous dates', 'Previous received collections (PHP)'])
    for row in receipt_rows(report):
        receipts.append([row['period'], row['value'], row['previous_period'], row['previous']])
    payments = workbook.create_sheet('Payment Breakdown')
    payments.append(['Section', 'Label', 'Booking count', 'Amount (PHP)'])
    for row in report['insights']['payment_progress']:
        payments.append(['Payment progress', row['label'], row['count']])
    for row in report['insights']['collections']['methods']:
        payments.append(['Receipt-period method', row['label'], None, row['value']])
    payments.append(['Reconciliation', 'Payable booking value', None, report['insights']['payable_value']])
    payments.append(['Reconciliation', 'Verified credit, capped per booking', None, report['insights']['credited_capped']])
    payments.append(['Reconciliation', 'Remaining balance', None, report['insights']['remaining_balance']])
    customers = workbook.create_sheet('Customer Breakdown')
    customers.append(['Section', 'Label', 'Count'])
    for group in ('types', 'ratings'):
        for row in report['insights']['customers'][group]:
            customers.append([group.title(), row['label'], row['count']])


def style_workbook(workbook):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.dimensions
        sheet.print_title_rows = '1:1'
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
        sheet.page_margins = PageMargins(left=.3, right=.3, top=.5, bottom=.5, header=.2, footer=.2)
        sheet.oddHeader.center.text = 'Balloorina | Reports & Analytics'
        sheet.oddFooter.right.text = 'Page &P of &N'
        sheet.row_dimensions[1].height = 30
        for row in sheet:
            sheet.row_dimensions[row[0].row].height = 30 if row[0].row > 1 else 32
            for cell in row:
                # Export user text as text, including strings beginning with '='.
                if isinstance(cell.value, str): cell.data_type = 's'
                header = cell.row == 1
                cell.font = Font(name='Calibri', size=11, bold=header, color='FFFFFF' if header else '262626')
                cell.fill = PatternFill('solid', fgColor='202020' if header else 'F6F6F6' if cell.row % 2 == 0 else 'FFFFFF')
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.border = Border(bottom=Side(style='hair', color='E6E6E6'))
                if isinstance(cell.value, (float, Decimal)):
                    cell.number_format = '#,##0.00;[Red](#,##0.00)'
                elif isinstance(cell.value, int): cell.number_format = '#,##0'
        for column in sheet.columns:
            width = min(46, max(16, max(len(str(cell.value or '')) for cell in column) + 3))
            sheet.column_dimensions[column[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            lines = max((ceil(len(str(cell.value or '')) / max(12, sheet.column_dimensions[cell.column_letter].width - 2)) for cell in row), default=1)
            sheet.row_dimensions[row[0].row].height = min(110, max(30, lines * 15 + 8))
        if sheet.title == 'Report Definitions':
            sheet.column_dimensions['A'].width = 26; sheet.column_dimensions['B'].width = 110
            for row in range(1, sheet.max_row + 1): sheet.row_dimensions[row].height = 45
