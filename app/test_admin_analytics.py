"""Cohort, calendar, payment and export parity checks for admin reports."""
import io
from datetime import date, datetime, timedelta, timezone as utc
from decimal import Decimal
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from .analytics_reporting import build_report, resolve_filters, _buckets
from .models import Booking, Payment, User


@override_settings(REQUEST_RATE_LIMITING_ENABLED=False)
class AdminAnalyticsTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='analytics-admin', role='admin', email_verified=True)
        self.customer = User.objects.create_user(username='analytics-customer', role='customer', email_verified=True)
        self.other = User.objects.create_user(username='analytics-other', role='customer', email_verified=True)
        self.client.force_login(self.admin)
        self.clock = self.enterContext(patch('app.analytics_reporting.timezone.now',
            return_value=datetime(2026, 10, 26, 4, tzinfo=utc.utc)))
        self.url = reverse('admin_analytics')

    def booking(self, day='2026-10-26', status='pending', value='1000.00', **kwargs):
        return Booking.objects.create(user=kwargs.pop('user', self.customer), event_date=day,
            event_type=kwargs.pop('event_type', 'Birthday'), package_type=kwargs.pop('package_type', 'Package 1-A + Entrance A'),
            event_location='Test venue', total_price=value, status=status, **kwargs)

    def payment(self, booking, value, status='verified', applied=True):
        return Payment.objects.create(booking=booking, amount=value, payment_status=status,
            applied_to_booking=applied, transaction_ref=f'ANALYTICS-{Payment.objects.count()}')

    def report(self, **params):
        return build_report(resolve_filters({'filter_preset': 'today', **params}))

    def test_all_metrics_share_cohort_and_payment_join_does_not_duplicate_bookings(self):
        done = self.booking(status='completed', value='1200.25')
        self.payment(done, '300.15'); self.payment(done, '200.10')
        self.payment(done, '100', status='pending'); self.payment(done, '500', status='rejected')
        self.payment(done, '900', applied=False)
        self.booking(status='cancelled', user=self.other)
        self.booking(status='cancel_requested')
        self.booking(day='2026-11-01', status='completed', value='9999')
        report = self.report()
        self.assertEqual(report['total_bookings'], 3)
        self.assertEqual(report['completed_value'], Decimal('1200.25'))
        self.assertEqual(report['verified_collections'], Decimal('500.25'))
        self.assertEqual(report['completion_rate'], 33.3)
        self.assertEqual(report['cancellation_rate'], 33.3)
        self.assertEqual(report['customer_count'], 2)
        self.assertEqual(sum(row['count'] for row in report['statuses']), 3)
        self.assertEqual(sum(row['count'] for row in report['event_rows']), 3)
        self.assertEqual(sum(row['count'] for row in report['setup_rows']), 3)
        self.assertEqual(sum(report['charts']['trend']['series']['collections']['current']), 500.25)
        self.assertIn('Cancel Requested', [row['label'] for row in report['statuses']])

    def test_event_and_submission_dates_use_philippine_midnight(self):
        before = self.booking(day='2026-12-10')
        beginning = self.booking(day='2026-12-10')
        last = self.booking(day='2026-12-10')
        after = self.booking(day='2026-12-10')
        for booking, stamp in [(before, '2026-10-25T15:59:59+00:00'), (beginning, '2026-10-25T16:00:00+00:00'),
                (last, '2026-10-26T15:59:59+00:00'), (after, '2026-10-26T16:00:00+00:00')]:
            Booking.objects.filter(pk=booking.pk).update(created_at=datetime.fromisoformat(stamp))
        self.assertEqual(self.report()['total_bookings'], 0)
        report = self.report(date_basis='created')
        self.assertEqual(report['total_bookings'], 2)
        self.assertEqual(report['previous_count'], 1)
        self.assertEqual(report['charts']['trend']['series']['bookings']['current'], [2])

    def test_today_is_server_philippine_date_not_utc_date(self):
        self.clock.return_value = datetime(2026, 10, 25, 16, 1, tzinfo=utc.utc)
        filters = resolve_filters({'filter_preset': 'today'})
        self.assertEqual(filters.start, date(2026, 10, 26))
        self.assertEqual(filters.previous_end, date(2026, 10, 25))

    def test_month_to_date_compares_matching_previous_month_days(self):
        self.booking(day='2026-10-26'); self.booking(day='2026-10-27')
        self.booking(day='2026-09-26'); self.booking(day='2026-09-27')
        report = self.report(filter_preset='month')
        self.assertEqual(report['filters'].end, date(2026, 10, 26))
        self.assertEqual(report['filters'].previous_start, date(2026, 9, 1))
        self.assertEqual(report['filters'].previous_end, date(2026, 9, 26))
        self.assertEqual((report['total_bookings'], report['previous_count']), (1, 1))

    def test_shorter_previous_month_has_no_invented_comparison_points(self):
        self.clock.return_value = datetime(2026, 3, 31, 4, tzinfo=utc.utc)
        report = self.report(filter_preset='month')
        self.assertEqual(report['filters'].previous_end, date(2026, 2, 28))
        self.assertEqual(report['charts']['trend']['series']['bookings']['previous'][-3:], [None] * 3)

    def test_custom_and_seven_day_comparison_are_adjacent_equal_windows(self):
        filters = resolve_filters({'start_date': '2026-10-20', 'end_date': '2026-10-26'})
        self.assertEqual((filters.previous_start, filters.previous_end), (date(2026, 10, 13), date(2026, 10, 19)))
        self.assertEqual(resolve_filters({'filter_preset': '7days'}).start, filters.start)

    def test_all_time_includes_future_and_early_bookings_without_fake_comparison(self):
        self.booking(day='1999-12-01'); self.booking(day='2027-02-03')
        report = self.report(filter_preset='all')
        self.assertEqual(report['total_bookings'], 2)
        self.assertEqual(report['filters'].start, date(1999, 12, 1))
        self.assertEqual(report['filters'].end, date(2027, 2, 3))
        self.assertIsNone(report['previous_count'])
        self.assertFalse(report['charts']['trend']['has_comparison'])
        self.assertLessEqual(len(report['charts']['trend']['labels']), 49)

    def test_event_type_filters_every_kpi_graph_and_all_time_bounds(self):
        self.booking(status='completed'); self.booking(event_type='Wedding', day='2030-01-01')
        report = self.report(filter_preset='all', event_type='Birthday')
        self.assertEqual(report['total_bookings'], 1)
        self.assertEqual(report['filters'].end, date(2026, 10, 26))
        self.assertEqual([row['label'] for row in report['event_rows']], ['Birthday'])

    def test_zero_baseline_does_not_claim_one_hundred_percent_growth(self):
        self.booking()
        delta = self.report()['comparison']['count']
        self.assertIsNone(delta['percent'])
        self.assertIn('no % baseline', delta['text'])
        self.booking(day='2026-10-25')
        self.assertEqual(self.report()['comparison']['count']['percent'], 0)
        self.booking()
        self.assertEqual(self.report()['comparison']['count']['percent'], 100)

    def test_empty_report_is_zero_and_never_counts_unrelated_accounts(self):
        report = self.report()
        self.assertEqual(report['customer_count'], 0)
        self.assertEqual(report['verified_collections'], 0)
        self.assertEqual(report['completion_rate'], 0)
        self.assertEqual(report['cancellation_rate'], 0)
        self.assertEqual(report['statuses'], [])

    def test_ranked_groups_preserve_counts_and_amounts_even_with_similar_event_name(self):
        for index, name in enumerate(['Other selections', 'Wedding', 'Birthday', 'A', 'B', 'C', 'D', 'E']):
            self.booking(event_type=name, status='completed', value='10.01', package_type=f'Setup {index} + Stand A')
        report = self.report()
        self.assertEqual(sum(row['count'] for row in report['event_rows']), 8)
        self.assertEqual(sum(row['value'] for row in report['event_rows']), Decimal('80.08'))
        self.assertEqual(sum(row['count'] for row in report['setup_rows']), 8)
        self.assertEqual(len(report['setup_rows']), 7)

    def test_interval_aggregation_preserves_totals_at_range_edges(self):
        for start, end in [('2026-01-01','2026-01-31'), ('2026-01-01','2026-04-30'),
                ('2025-06-15','2026-07-10'), ('2020-06-15','2026-07-10')]:
            with self.subTest(start=start):
                Booking.objects.all().delete()
                self.booking(day=start); self.booking(day=end)
                report = self.report(filter_preset='', start_date=start, end_date=end)
                self.assertEqual(sum(report['charts']['trend']['series']['bookings']['current']), 2)
                self.assertEqual(report['total_bookings'], 2)
        self.assertLessEqual(len(_buckets(date.min, date.max)[1]), 49)

    def test_aggregation_query_count_does_not_grow_with_bookings(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext
        filters = resolve_filters({'filter_preset':'today'})
        self.booking()
        with CaptureQueriesContext(connection) as initial:
            build_report(filters)
        for _ in range(8): self.booking()
        with CaptureQueriesContext(connection) as expanded:
            build_report(filters)
        self.assertEqual(len(initial), len(expanded))
        self.assertLessEqual(len(expanded), 8)

    def test_bad_filters_return_validation_errors_not_silently_swapped_periods(self):
        for params in [{'date_basis':'payment'}, {'event_type':'Missing'}, {'filter_preset':'invalid'},
                {'start_date':'2026-02-30','end_date':'2026-03-01'}, {'start_date':'2026-10-01'},
                {'start_date':'2026-10-27','end_date':'2026-10-26'}]:
            with self.subTest(params=params):
                with self.assertRaises(ValidationError): resolve_filters(params)
                response = self.client.get(self.url, params, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
                self.assertEqual(response.status_code, 400)
                self.assertTrue(response.json()['error'])
                self.assertIn('no-store', response['Cache-Control'])

    def test_html_and_async_fragment_display_same_numbers_and_safe_labels(self):
        self.booking(status='completed', event_type='<script>alert(1)</script>')
        params = {'filter_preset':'today'}
        full = self.client.get(self.url, params)
        self.assertEqual(full.status_code, 200)
        fragment = self.client.get(self.url, params, HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(fragment.status_code, 200)
        self.assertIn('₱1,000.00', fragment.json()['html'])
        self.assertNotIn('<script>alert(1)</script>', fragment.json()['html'])
        self.assertNotContains(full, 'Recent Bookings')
        self.assertNotContains(full, 'retentionChart')
        self.assertIn('no-store', full['Cache-Control'])

    def test_export_dates_are_frozen_and_match_report_cohort(self):
        self.booking(day='2026-10-26', status='completed', value='2345.67')
        report = self.report()
        params = parse_qs(urlsplit(report['export_url']).query)
        self.assertNotIn('filter_preset', params)
        self.clock.return_value += timedelta(days=1)
        with patch('app.views.admin.pisa.CreatePDF') as pdf:
            pdf.return_value.err = 0
            self.client.get(report['export_url'])
            html = pdf.call_args.args[0]
        self.assertIn('PHP 2,345.67', html)
        self.assertIn('Oct 26, 2026', html)
        self.assertIn('Event date', html)

    def test_excel_matches_event_date_cohort_and_only_includes_credited_payments(self):
        import openpyxl
        selected = self.booking(status='completed')
        Booking.objects.filter(pk=selected.pk).update(created_at=timezone.now()-timedelta(days=40))
        self.payment(selected, '100'); self.payment(selected, '200', status='pending')
        self.booking(day='2026-12-01')
        response = self.client.get(reverse('admin_analytics_export_excel'), {'filter_preset':'today'})
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(book['Bookings Summary'].max_row, 2)
        self.assertEqual(book['Verified Collections'].max_row, 2)
        self.assertEqual(book['Verified Collections']['C2'].value, 100)

    def test_reading_report_does_not_mutate_bookings(self):
        booking = self.booking(status='pending_payment')
        self.client.get(self.url)
        booking.refresh_from_db()
        self.assertEqual(booking.status, 'pending_payment')

    def test_exports_reject_reversed_dates_and_pdf_failures_return_errors(self):
        for name in ('admin_analytics_export_pdf', 'admin_analytics_export_excel'):
            response = self.client.get(reverse(name), {'start_date':'2026-10-27','end_date':'2026-10-26'})
            self.assertEqual(response.status_code, 400)
        with patch('app.views.admin.pisa.CreatePDF', side_effect=RuntimeError('Internal renderer detail')):
            response = self.client.get(reverse('admin_analytics_export_pdf'))
            self.assertEqual(response.status_code, 500)
            self.assertNotIn('Internal renderer detail', response.content.decode())
        with patch('app.views.admin.pisa.CreatePDF') as pdf:
            pdf.return_value.err = 1
            self.assertEqual(self.client.get(reverse('admin_analytics_export_pdf')).status_code, 500)

    def test_admin_only_get_endpoints(self):
        urls = [self.url, reverse('admin_analytics_export_pdf'), reverse('admin_analytics_export_excel')]
        for url in urls:
            self.assertEqual(self.client.post(url).status_code, 405)
        self.client.force_login(self.customer)
        for url in urls: self.assertEqual(self.client.get(url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
